# -*- coding: utf-8 -*-
"""Bygger vejledning/grupper-fredag.docx — samme liste som PDF'en, i Word.

Navnene hentes fra claude/byg_grupper.py, saa de to filer ikke kan komme til
at sige hver sit. Kraever python-docx, som her kun ligger i python3.13:

    python3.13 claude/byg_grupper_word.py
"""
import os
import sys

from docx import Document
from docx.shared import Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import byg_grupper as G

UD = os.path.join(G.ROD, 'vejledning', 'grupper-fredag.docx')



def _kort(p, tekst, stoerrelse, fed=False):
    k = p.add_run(tekst)
    k.font.size = Pt(stoerrelse)
    k.bold = fed
    return k


dok = Document()
stil = dok.styles['Normal']
stil.font.name = 'Times New Roman'
stil.font.size = Pt(12)

p = dok.add_paragraph()
_kort(p, G.TITEL, 18, fed=True)
p.paragraph_format.space_after = Pt(2)

p = dok.add_paragraph()
_kort(p, G.UNDER, 12)
p.paragraph_format.space_after = Pt(18)

# List Number er Words egen nummerering. Den taeller selv, saa numrene
# retter sig, hvis en gruppe bliver slettet eller flyttet.
for navne in G.GRUPPER:
    p = dok.add_paragraph(style='List Number')
    _kort(p, ', '.join(navne), 12)
    p.paragraph_format.space_after = Pt(6)

p = dok.add_paragraph()
p.paragraph_format.space_before = Pt(18)
_kort(p, f'{len(G.GRUPPER)} grupper · {G.I_ALT} unger.', 11)

dok.save(UD)
print(f'skrevet:  {UD}')
print(f'{len(G.GRUPPER)} grupper · {G.I_ALT} unger')
