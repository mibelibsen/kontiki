# -*- coding: utf-8 -*-
"""Bygger vejledning/staaende-til-liggende-video.pdf og .docx

Hvordan man faar en staaende video til at blive liggende. Der er to helt
forskellige problemer, som ligner hinanden, og vejledningen skiller dem ad:

  1. Billedet ligger paa siden        -> den skal bare roteres
  2. Billedet staar rigtigt, men er   -> den skal have et nyt format, og man
     hoejt og smalt                      skal vaelge mellem at beskaere og
                                         at faa bjaelker i siderne

    python3 claude/byg_videoformat.py          # PDF
    python3.13 claude/byg_videoformat.py --word

Teksten staar ét sted (DEL og HUSK), saa de to filer ikke kan sige hver sit.
"""
import html
import os
import re
import subprocess
import sys

ROD = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NAVN = 'staaende-til-liggende-video'
UD_HTML = os.path.join(ROD, 'vejledning', NAVN + '.html')
UD_PDF = os.path.join(ROD, 'vejledning', NAVN + '.pdf')
UD_PNG = os.path.join(ROD, 'vejledning', '.' + NAVN + '-figur.png')
UD_DOCX = os.path.join(ROD, 'vejledning', NAVN + '.docx')

TITEL = 'Fra stående til liggende video'
UNDER = ('Der er to forskellige problemer, som ligner hinanden. '
         'Start med at afspille videoen og find ud af, hvilket et du har.')

# (overskrift, hvornaar, [(maskine, [trin])], note)
DEL = [
    ('1. Billedet ligger på siden',
     'Personerne i videoen ligger ned, som om du skal dreje hovedet for at '
     'se den. Så er videoen optaget rigtigt — den er bare gemt forkert '
     'vendt. Den skal roteres en kvart omgang.',
     [
      ('iPad eller iPhone', [
         'Åbn videoen i <b>Billeder</b>, tryk <b>Rediger</b>.',
         'Tryk på <b>beskær-ikonet</b> nederst.',
         'Tryk på <b>rotér-ikonet</b> øverst til venstre — det '
         '<b>firkantede</b> med pilen. Tryk igen, til billedet står rigtigt.',
         'Tryk <b>Færdig</b>.',
      ]),
      ('Mac', [
         'Åbn videoen i <b>QuickTime Player</b>.',
         'Vælg <b>Rediger → Rotér til venstre</b> eller <b>Rotér til '
         'højre</b>.',
         'Luk vinduet, og tryk <b>Arkivér</b>.',
      ]),
      ('Windows', [
         'Åbn <b>Clipchamp</b> og træk videoen ned på tidslinjen.',
         'Klik på klippet, så det får en grøn ramme.',
         'Brug <b>rotér</b> i værktøjslinjen ved siden af '
         'forhåndsvisningen.',
         'Tryk <b>Eksportér</b> og vælg <b>1080p</b>.',
      ]),
     ],
     'Er det det her, du har, så er du færdig. Resten af arket handler om '
     'det andet problem.'),

    ('2. Billedet står rigtigt, men er højt og smalt',
     'Videoen er optaget med telefonen på højkant. Den skal nu være bred '
     '(16:9), for eksempel til at vise på et lærred. <b>Her skal du '
     'vælge</b>: et højt billede kan ikke blive bredt, uden at der enten '
     'ryger noget af toppen og bunden, eller kommer tomme bjælker i '
     'siderne. Se tegningen øverst.',
     [
      ('Windows — her kan du vælge selv', [
         'Åbn <b>Clipchamp</b>, lav et nyt projekt, og træk videoen ned på '
         'tidslinjen.',
         'Klik på <b>formatknappen</b> lige under Eksportér øverst til '
         'højre, og vælg <b>16:9</b>.',
         'Klik på klippet, så det får en grøn ramme.',
         'Vælg <b>Udfyld</b> i værktøjslinjen, hvis billedet skal fylde det '
         'hele — så beskæres top og bund. Vælg <b>Tilpas</b>, hvis intet må '
         'ryge — så kommer der bjælker i siderne.',
         'Tryk <b>Eksportér</b> og vælg <b>1080p</b>.',
      ]),
      ('iPad eller iPhone — beskærer altid', [
         'Åbn videoen i <b>Billeder</b>, tryk <b>Rediger</b>.',
         'Tryk på <b>beskær-ikonet</b> nederst.',
         'Tryk på <b>formatknappen</b> øverst til højre og vælg '
         '<b>16:9</b>.',
         'Træk i billedet, så det vigtigste er inden for rammen.',
         'Tryk <b>Færdig</b>.',
      ]),
      ('Mac', [
         'Åbn <b>iMovie</b> og lav et nyt projekt.',
         'Træk videoen ind på tidslinjen og klik på den.',
         'Vælg <b>beskæringsværktøjet</b> over forhåndsvisningen.',
         'Vælg <b>Beskær til fuld</b> (beskærer) eller <b>Tilpas</b> '
         '(bjælker).',
         'Vælg <b>Del → Arkivér fil</b> og gem i 1080p.',
      ]),
     ],
     ''),
]

HUSK = [
    'Gem en kopi af den oprindelige video, før du beskærer. Det, der er '
    'beskåret væk, kommer ikke igen.',
    'Skal videoen vises på et lærred eller en projektor, er <b>16:9</b> '
    'det rigtige format.',
    'Lad være med at søge efter »convert video online«. De fleste af de '
    'sider vil have dig til at <b>uploade</b> filen til en server, du ikke '
    'ved, hvor står. Alle metoderne her foregår på din egen maskine.',
]

# ---------------------------------------------------------------- figuren
# Tre rammer: videoen som den er, og de to maader den kan blive bred paa.
# Alle maal er regnet, ikke sat paa oejemaal: de to brede rammer er lige
# store, de staar symmetrisk om midten, og personen har samme hoejde/bredde
# -forhold alle tre steder.
B_HOEJ, H_HOEJ = 62, 112
B_BRED, H_BRED = 160, 90
GAB1, GAB2 = 52, 42
W = 2 * 40 + B_HOEJ + GAB1 + B_BRED + GAB2 + B_BRED
H = 168
Y = 16
X_HOEJ = 40
X_A = X_HOEJ + B_HOEJ + GAB1
X_B = X_A + B_BRED + GAB2
assert X_B + B_BRED + 40 == W, 'rammerne står ikke symmetrisk'
Y_BRED = Y + (H_HOEJ - H_BRED) / 2
MIDT_Y = Y + H_HOEJ / 2

INK, GRAA, LYS = '#1a2233', '#586074', '#d8dee9'


def _person(cx, top, hoejde, ident):
    """En figur paa hoejkant: hoved og krop. `hoejde` er hele figuren, og
    hovedet er en femtedel af den - samme forhold hver gang, saa de tre
    tegninger viser det samme motiv i forskellig stoerrelse."""
    r = hoejde / 10
    krop_top = top + 2.6 * r
    return (f'<circle cx="{cx:.1f}" cy="{top + r:.1f}" r="{r:.1f}" '
            f'fill="{GRAA}"/>'
            f'<rect x="{cx - 1.7 * r:.1f}" y="{krop_top:.1f}" '
            f'width="{3.4 * r:.1f}" height="{top + hoejde - krop_top:.1f}" '
            f'rx="{r:.1f}" fill="{GRAA}"/>')


BASIS = Y + H_HOEJ + 16          # samme tekstlinje under alle tre rammer


def _ramme(x, y, b, h, ident, indhold, maerke, forklaring):
    return (f'<clipPath id="{ident}"><rect x="{x}" y="{y}" width="{b}" '
            f'height="{h}" rx="3"/></clipPath>'
            f'<rect x="{x}" y="{y}" width="{b}" height="{h}" rx="3" '
            f'fill="#fff" stroke="{INK}" stroke-width="1.2"/>'
            f'<g clip-path="url(#{ident})">{indhold}</g>'
            f'<rect x="{x}" y="{y}" width="{b}" height="{h}" rx="3" '
            f'fill="none" stroke="{INK}" stroke-width="1.2"/>'
            f'<text x="{x + b / 2:.1f}" y="{BASIS:.1f}" '
            f'text-anchor="middle" font-size="12.5" font-weight="700" '
            f'fill="{INK}">{maerke}</text>'
            f'<text x="{x + b / 2:.1f}" y="{BASIS + 15:.1f}" '
            f'text-anchor="middle" font-size="11" fill="{GRAA}">'
            f'{forklaring}</text>')


def _pil(x1, x2):
    return (f'<line x1="{x1:.1f}" y1="{MIDT_Y:.1f}" x2="{x2 - 9:.1f}" '
            f'y2="{MIDT_Y:.1f}" stroke="{INK}" stroke-width="1.5"/>'
            f'<path d="M{x2:.1f},{MIDT_Y:.1f} l-9,-5 v10 z" fill="{INK}"/>')


# Som den er: figuren fylder rammen paa hoejkant.
_FIG_HOEJ = _person(X_HOEJ + B_HOEJ / 2, Y + 12, H_HOEJ - 24, 'a')
# Beskaaret: for at fylde en bred ramme ud skal et hoejt billede forstoerres,
# til det er bredt nok - og saa falder top og bund uden for rammen. Figuren
# er derfor 1,7 gange rammens hoejde og placeret, saa hovedet lige er skaaret
# til i toppen og kroppen loeber ud forneden.
_A_HOEJ = H_BRED * 1.7
_FIG_A = _person(X_A + B_BRED / 2, Y_BRED - _A_HOEJ / 10 + 2, _A_HOEJ, 'b')
# Bjaelker: figuren er skaleret ned, saa den er inden for den lave ramme,
# og der er graa felter i siderne.
_SKALA = (H_BRED - 16) / (H_HOEJ - 24)
_BJ_B = B_HOEJ * _SKALA
_FIG_B = (f'<rect x="{X_B}" y="{Y_BRED}" width="{(B_BRED - _BJ_B) / 2:.1f}" '
          f'height="{H_BRED}" fill="{LYS}"/>'
          f'<rect x="{X_B + (B_BRED + _BJ_B) / 2:.1f}" y="{Y_BRED}" '
          f'width="{(B_BRED - _BJ_B) / 2:.1f}" height="{H_BRED}" '
          f'fill="{LYS}"/>'
          + _person(X_B + B_BRED / 2, Y_BRED + 8, H_BRED - 16, 'c'))

FIGUR = (
    f'<svg viewBox="0 0 {W} {H}" width="100%" style="max-width:{W}px" '
    f'xmlns="http://www.w3.org/2000/svg" role="img" '
    f'aria-label="En høj video kan blive bred på to måder: ved at beskære '
    f'top og bund væk, eller ved at få grå bjælker i siderne." '
    f'font-family="Helvetica,Arial,sans-serif">'
    + _ramme(X_HOEJ, Y, B_HOEJ, H_HOEJ, 'r1', _FIG_HOEJ,
             'Sådan er den', 'høj og smal')
    + _pil(X_HOEJ + B_HOEJ + 12, X_A - 12)
    + _ramme(X_A, Y_BRED, B_BRED, H_BRED, 'r2', _FIG_A,
             'Beskær', 'top og bund ryger')
    + _ramme(X_B, Y_BRED, B_BRED, H_BRED, 'r3', _FIG_B,
             'Bjælker', 'intet ryger')
    + f'<text x="{(X_A + B_BRED + X_B) / 2:.1f}" y="{MIDT_Y + 4:.1f}" '
      f'text-anchor="middle" font-size="12" fill="{GRAA}">eller</text>'
    + '</svg>')

CSS = '''
body{margin:0;font-family:Georgia,'Times New Roman',serif;color:#000;
font-size:10.5pt;line-height:1.38}
h1{font-size:17pt;margin:0 0 2mm}
h2{font-size:12.5pt;margin:4mm 0 1mm}
h3{font-size:10.5pt;margin:2.5mm 0 1mm}
.under{margin:0 0 3mm}
.figur{margin:0 0 4mm;text-align:center}
.hvornaar{margin:0 0 1.5mm;color:#333}
ol{margin:0;padding-left:9mm}
ol li{margin:0 0 1mm;padding-left:2mm}
.noteline{margin:1.5mm 0 0;color:#333}
ul{margin:0;padding-left:9mm}
ul li{margin:0 0 1.5mm;padding-left:2mm}
h2.husk{margin-top:4mm}
@page{size:A4;margin:13mm}
'''


def _afsnit_html():
    ud = ''
    for titel, hvornaar, maskiner, note in DEL:
        ud += f'<h2>{html.escape(titel)}</h2><p class="hvornaar">{hvornaar}</p>'
        for navn, trin in maskiner:
            ud += (f'<h3>{html.escape(navn)}</h3><ol>'
                   + ''.join(f'<li>{t}</li>' for t in trin) + '</ol>')
        if note:
            ud += f'<p class="noteline">{note}</p>'
    return ud


DOK = ('<!DOCTYPE html><html lang="da"><head><meta charset="UTF-8">'
       f'<title>{html.escape(TITEL)}</title><style>' + CSS
       + '</style></head><body>'
       f'<h1>{html.escape(TITEL)}</h1>'
       f'<p class="under">{html.escape(UNDER)}</p>'
       f'<div class="figur">{FIGUR}</div>'
       + _afsnit_html()
       + '<h2 class="husk">Husk</h2><ul>'
       + ''.join(f'<li>{t}</li>' for t in HUSK)
       + '</ul></body></html>')


def _skriv(p, s, fed_alt=False):
    """Skriver teksten og saetter <b>...</b> fed undervejs, saa Word-udgaven
    fremhaever de samme ord som PDF'en."""
    for stump in re.split(r'(<b>.*?</b>)', s):
        if not stump:
            continue
        fed = stump.startswith('<b>')
        k = p.add_run(stump[3:-4] if fed else stump)
        k.bold = fed or fed_alt
    return p


def byg_pdf():
    os.makedirs(os.path.dirname(UD_HTML), exist_ok=True)
    open(UD_HTML, 'w', encoding='utf-8').write(DOK)
    subprocess.run(['node', os.path.join(ROD, 'claude', 'html_til_pdf.mjs'),
                    UD_HTML, UD_PDF], check=True, stdout=subprocess.DEVNULL)
    os.remove(UD_HTML)
    print(f'skrevet:  {UD_PDF}')


def byg_figur_png():
    h = ('<!DOCTYPE html><meta charset="UTF-8">'
         '<style>body{margin:0}svg{display:block}</style>' + FIGUR)
    sti = UD_HTML.replace('.html', '-figur.html')
    open(sti, 'w', encoding='utf-8').write(h)
    subprocess.run(['node', os.path.join(ROD, 'claude', 'svg_til_png.mjs'),
                    sti, UD_PNG, str(W), str(H)], check=True)
    os.remove(sti)
    return UD_PNG


def byg_word():
    from docx import Document
    from docx.shared import Cm, Pt

    png = byg_figur_png()
    dok = Document()
    for s in dok.sections:
        s.top_margin = s.bottom_margin = Cm(1.3)
        s.left_margin = s.right_margin = Cm(1.6)
    dok.styles['Normal'].font.name = 'Times New Roman'
    dok.styles['Normal'].font.size = Pt(10.5)

    p = dok.add_paragraph()
    k = p.add_run(TITEL)
    k.bold = True
    k.font.size = Pt(17)
    dok.add_paragraph(UNDER).paragraph_format.space_after = Pt(6)
    dok.add_picture(png, width=Cm(14))
    dok.paragraphs[-1].alignment = 1
    dok.paragraphs[-1].paragraph_format.space_after = Pt(8)

    # Numrene skrives som tekst. Words egen List Number deler én taeller
    # mellem alle lister i dokumentet, saa anden liste ville begynde ved 5.
    for titel, hvornaar, maskiner, note in DEL:
        h = dok.add_paragraph()
        r = h.add_run(titel)
        r.bold = True
        r.font.size = Pt(12.5)
        h.paragraph_format.space_before = Pt(9)
        _skriv(dok.add_paragraph(), hvornaar)
        for navn, trin in maskiner:
            m = dok.add_paragraph()
            m.paragraph_format.space_before = Pt(5)
            m.paragraph_format.space_after = Pt(1)
            mr = m.add_run(navn)
            mr.bold = True
            for nr, t in enumerate(trin, 1):
                a = dok.add_paragraph()
                a.paragraph_format.left_indent = Cm(1.1)
                a.paragraph_format.first_line_indent = Cm(-0.7)
                a.paragraph_format.space_after = Pt(1)
                a.add_run(f'{nr}.\t')
                _skriv(a, t)
        if note:
            n = dok.add_paragraph()
            n.paragraph_format.space_before = Pt(4)
            _skriv(n, note)

    h = dok.add_paragraph()
    r = h.add_run('Husk')
    r.bold = True
    r.font.size = Pt(12.5)
    h.paragraph_format.space_before = Pt(10)
    for t in HUSK:
        _skriv(dok.add_paragraph(style='List Bullet'), t)

    dok.save(UD_DOCX)
    os.remove(png)
    print(f'skrevet:  {UD_DOCX}')


if __name__ == '__main__':
    byg_pdf()
    if '--word' in sys.argv:
        byg_word()
    print(f'{len(DEL)} problemer · '
          f'{sum(len(t) for _, _, m, _ in DEL for _, t in m)} trin i alt')
