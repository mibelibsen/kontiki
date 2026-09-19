#!/usr/bin/env python3
"""Bygger en vejledning som Word-fil ud af den HTML, PDF'en gengives fra.

Ingen tekst skrives af i hånden: scriptet læser overskrifter, afsnit, tabeller
og figurer ud af HTML'en, så Word-filen og PDF'en ikke kan komme til at sige
noget forskelligt.

    python3 claude/byg_vejledning_docx.py <kilde.html> <ud.docx>

Figurerne er SVG i HTML'en. De gengives som PNG med Chromium (Playwright) og
lægges ind samme sted i dokumentet, som de står på siden.
"""
import json
import math
import os
import re
import subprocess
import sys
from html.parser import HTMLParser

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

KILDE, UD = sys.argv[1], sys.argv[2]
ROD = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ARBEJDE = os.path.dirname(os.path.abspath(KILDE))

SKRIFT = 'Calibri'
INK = RGBColor(0x1A, 0x22, 0x33)
MUTED = RGBColor(0x58, 0x60, 0x74)
ACCENT = RGBColor(0x1F, 0x6F, 0xD6)
LINJE = 'D3DAE7'
HOVED_FYLD = 'F4F6FB'

# =============================================================== 1 · læs HTML

class Laeser(HTMLParser):
    """Trækker blokke ud af vejledningens HTML.

    Resultatet er en liste af (slags, indhold):
      ('h1'|'h2'|'h3', tekststykker) · ('p', tekststykker) · ('figur', None)
      ('tabel', {'hoved': [...], 'raekker': [[...]], 'tal': [bool]})
    Tekststykker er (tekst, fed, kursiv), så fremhævninger overlever.
    """

    SPRING_OVER = {'header', 'footer', 'button', 'style', 'script', 'nav'}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blokke = []
        self.dybde_spring = 0
        self.i_svg = 0
        self.aktiv = None          # slags for det afsnit vi samler på
        self.dele = []
        self.fed = self.kursiv = 0
        self.tabel = None
        self.celle = None
        self.i_hoved = False

    # -- små hjælpere ----------------------------------------------------
    def _luk_afsnit(self):
        if self.aktiv and any(t.strip() for t, _, _ in self.dele):
            self.blokke.append((self.aktiv, self._ryd(self.dele)))
        self.aktiv, self.dele = None, []

    @staticmethod
    def _ryd(dele):
        ud = []
        for tekst, f, k in dele:
            tekst = re.sub(r'\s+', ' ', tekst)
            if not tekst:
                continue
            if not ud and tekst.startswith(' '):
                tekst = tekst.lstrip()
            if ud and ud[-1][1] == f and ud[-1][2] == k:
                ud[-1] = (ud[-1][0] + tekst, f, k)
            else:
                ud.append((tekst, f, k))
        if ud:
            ud[-1] = (ud[-1][0].rstrip(), ud[-1][1], ud[-1][2])
        return [d for d in ud if d[0]]

    # -- parserens tilbagekald -------------------------------------------
    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == 'svg':
            self.i_svg += 1
            if self.i_svg == 1:
                self._luk_afsnit()
                self.blokke.append(('figur', None))
            return
        if self.i_svg:
            return
        if self.dybde_spring or tag in self.SPRING_OVER:
            self.dybde_spring += 1
            return
        if tag in ('h1', 'h2', 'h3', 'p'):
            self._luk_afsnit()
            self.aktiv = tag
        elif tag == 'table':
            self._luk_afsnit()
            self.tabel = {'hoved': [], 'raekker': [], 'tal': []}
        elif tag == 'tr' and self.tabel is not None:
            if not self.i_hoved:
                self.tabel['raekker'].append([])
        elif tag in ('th', 'td') and self.tabel is not None:
            self.celle = []
            if tag == 'td' and len(self.tabel['raekker']) == 1:
                self.tabel['tal'].append('tal' in a.get('class', ''))
        elif tag == 'thead':
            self.i_hoved = True
        elif tag in ('b', 'strong'):
            self.fed += 1
        elif tag in ('i', 'em'):
            self.kursiv += 1
        elif tag == 'br':
            self.dele.append((' ', self.fed > 0, self.kursiv > 0))

    def handle_endtag(self, tag):
        if tag == 'svg':
            self.i_svg = max(0, self.i_svg - 1)
            return
        if self.i_svg:
            return
        if self.dybde_spring:
            self.dybde_spring -= 1
            return
        if tag in ('h1', 'h2', 'h3', 'p'):
            self._luk_afsnit()
        elif tag == 'thead':
            self.i_hoved = False
        elif tag in ('th', 'td') and self.tabel is not None:
            tekst = ' '.join(t for t, _, _ in self._ryd(self.celle or []))
            if tag == 'th':
                self.tabel['hoved'].append(tekst)
            else:
                self.tabel['raekker'][-1].append(tekst)
            self.celle = None
        elif tag == 'table' and self.tabel is not None:
            self.tabel['raekker'] = [r for r in self.tabel['raekker'] if r]
            self.blokke.append(('tabel', self.tabel))
            self.tabel = None
        elif tag in ('b', 'strong'):
            self.fed = max(0, self.fed - 1)
        elif tag in ('i', 'em'):
            self.kursiv = max(0, self.kursiv - 1)

    def handle_data(self, data):
        if self.i_svg or self.dybde_spring:
            return
        stykke = (data, self.fed > 0, self.kursiv > 0)
        if self.celle is not None:
            self.celle.append(stykke)
        elif self.aktiv:
            self.dele.append(stykke)


html = open(KILDE, encoding='utf-8').read()
laeser = Laeser()
laeser.feed(html[html.index('<body'):])
BLOKKE = laeser.blokke

assert BLOKKE and BLOKKE[0][0] == 'h1', 'forventede en h1 først i vejledningen'
assert any(s == 'p' for s, _ in BLOKKE), 'fandt ingen brødtekst'
_pill = re.search(r'class="pill">([^<]+)<', html)
UNDERTITEL = _pill.group(1) if _pill else ''

# =============================================================== 2 · figuren

SVGER = re.findall(r'<svg.*?</svg>', html, re.S)
assert len(SVGER) == sum(1 for s, _ in BLOKKE if s == 'figur'), \
    'parseren fandt ikke lige så mange figurer, som HTML-filen har'

# Hver SVG gengives som PNG i tre gange opløsning, så stregerne holder i print.
JOBS = []
for nr, svg in enumerate(SVGER):
    b, h = (float(v) for v in
            re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', svg).groups())
    fast = re.sub(r'width="100%"', f'width="{b:.0f}"', svg, count=1)
    open(os.path.join(ARBEJDE, f'_figur{nr}.html'), 'w', encoding='utf-8').write(
        '<!doctype html><meta charset="utf-8">'
        '<body style="margin:0;background:#fff">' + fast + '</body>')
    JOBS.append([nr, math.ceil(b), math.ceil(h)])

subprocess.run(['node', '-e', f"""
import('/opt/node22/lib/node_modules/playwright/index.mjs').then(async (m) => {{
  const mappe = {json.dumps(os.path.abspath(ARBEJDE))};
  const b = await m.chromium.launch();
  const c = await b.newContext({{deviceScaleFactor: 3}});
  for (const [nr, bredde, hoejde] of {json.dumps(JOBS)}) {{
    const p = await c.newPage();
    await p.setViewportSize({{width: bredde, height: hoejde}});
    await p.goto('file://' + mappe + '/_figur' + nr + '.html');
    await p.waitForTimeout(200);
    await p.screenshot({{path: mappe + '/_vejl-figur' + nr + '.png'}});
    await p.close();
  }}
  await b.close();
}});"""], check=True)

FIGUR_KOE = [os.path.join(ARBEJDE, f'_vejl-figur{nr}.png') for nr, _, _ in JOBS]
for png in FIGUR_KOE:
    assert os.path.getsize(png) > 2000, f'{png} blev ikke gengivet'

# =============================================================== 3 · skriv docx

# Rækkefølgen af børn i <w:tcPr> er fastlagt i skemaet: tcW, tcBorders, shd,
# … , vAlign. Word nægter at åbne filen, hvis de står i en anden rækkefølge, så
# hvert element sættes ind foran sine efterfølgere i stedet for bare at hænges
# bagpå.
EFTER_TCBORDERS = ('w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection',
                   'w:tcFitText', 'w:vAlign', 'w:hideMark')
EFTER_SHD = EFTER_TCBORDERS[1:]


def _saet(celle, element, efterfoelgere):
    pr = celle._tc.get_or_add_tcPr()
    pr.insert_element_before(element, *efterfoelgere)


def kant(celle, farve=LINJE):
    ramme = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '6')
        b.set(qn('w:color'), farve)
        ramme.append(b)
    _saet(celle, ramme, EFTER_TCBORDERS)


def fyld(celle, farve):
    sk = OxmlElement('w:shd')
    sk.set(qn('w:val'), 'clear')
    sk.set(qn('w:fill'), farve)
    _saet(celle, sk, EFTER_SHD)


def skriv(afsnit, dele, storrelse=11, farve=INK, fed_alt=False):
    for tekst, f, k in dele:
        r = afsnit.add_run(tekst)
        r.font.name = SKRIFT
        r.font.size = Pt(storrelse)
        r.font.color.rgb = farve
        r.bold = f or fed_alt
        r.italic = k
    return afsnit


doc = Document()
sek = doc.sections[0]
sek.top_margin = sek.bottom_margin = Cm(2.0)
sek.left_margin = sek.right_margin = Cm(2.2)
BRED = sek.page_width - sek.left_margin - sek.right_margin

normal = doc.styles['Normal']
normal.font.name = SKRIFT
normal.font.size = Pt(11)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for slags, indhold in BLOKKE:
    if slags == 'h1':
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        skriv(p, indhold, storrelse=24, fed_alt=True)
        u = doc.add_paragraph()
        u.paragraph_format.space_after = Pt(14)
        skriv(u, [(UNDERTITEL, False, False)],
              storrelse=11, farve=ACCENT, fed_alt=True)
    elif slags in ('h2', 'h3'):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16 if slags == 'h2' else 10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        skriv(p, indhold, storrelse=14 if slags == 'h2' else 12, fed_alt=True)
    elif slags == 'p':
        skriv(doc.add_paragraph(), indhold)
    elif slags == 'figur':
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.add_run().add_picture(FIGUR_KOE.pop(0), width=min(BRED, Cm(15)))
    elif slags == 'tabel':
        kolonner = len(indhold['hoved'])
        t = doc.add_table(rows=1, cols=kolonner)
        t.autofit = False
        bredder = [int(BRED / kolonner)] * kolonner
        for i, navn in enumerate(indhold['hoved']):
            c = t.rows[0].cells[i]
            c.width = bredder[i]
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            fyld(c, HOVED_FYLD)
            kant(c)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            skriv(p, [(navn, False, False)], storrelse=10, farve=MUTED,
                  fed_alt=True)
        for raekke in indhold['raekker']:
            celler = t.add_row().cells
            for i, tekst in enumerate(raekke[:kolonner]):
                c = celler[i]
                c.width = bredder[i]
                kant(c)
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                if i < len(indhold['tal']) and indhold['tal'][i]:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                skriv(p, [(tekst, False, False)], storrelse=10)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

doc.save(UD)

print(f'skrevet:  {os.path.relpath(UD, ROD)}')
print(f'blokke:   {len(BLOKKE)} · '
      f"{sum(1 for s, _ in BLOKKE if s == 'tabel')} tabeller · "
      f"{sum(1 for s, _ in BLOKKE if s.startswith('h'))} overskrifter")
