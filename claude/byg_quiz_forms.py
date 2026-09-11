# -*- coding: utf-8 -*-
"""Bygger Forms-udgaven af statistik-testen ud fra quiz.json.

Microsoft Forms har ingen regnearksimport. Den importerer fra dokument, og
billeder kan ikke folge med - de saettes ind pr. sporgsmaal bagefter. Derfor
to filer: et Word-dokument til importen og en ren tekstfil til at kopiere fra,
hvis importen driller.

Koer:  python3 claude/byg_quiz_forms.py <mappe-med-quiz.json>
"""
import sys, os, json
import docx
from docx.shared import Pt, RGBColor

FIGDIR = sys.argv[1] if len(sys.argv) > 1 else 'quizfig'
SP = json.load(open(os.path.join(FIGDIR, 'quiz.json')))
assert len(SP) == 30, f'der skal vaere 30 spoergsmaal, ikke {len(SP)}'
BOGSTAV = 'ABCD'
DATO = '2026-09-11'

# ------------------------------------------------------------------- Word
d = docx.Document()
d.styles['Normal'].font.name = 'Arial'
d.styles['Normal'].font.size = Pt(11)
d.add_heading('Statistik · 9. klasse', 0)
for t, lille in (
        ('30 spørgsmål med fire svarmuligheder. Det rigtige svar er markeret '
         'med en stjerne (*) foran.', True),
        ('Billederne kan ikke følge med i en import. De ligger i mappen '
         'kahoot-billeder og er navngivet med spørgsmålets nummer — sæt dem ind '
         'i Forms bagefter.', True),
        ('Facit må ikke deles med ungerne.', True)):
    p = d.add_paragraph(t)
    if lille:
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.color.rgb = RGBColor(0x58, 0x60, 0x74)

for s in SP:
    p = d.add_paragraph()
    p.add_run(f'{s["nr"]}. {s["q"]}').bold = True
    for j, svar in enumerate(s['sv']):
        rigtigt = j == s['rigtig']
        linje = d.add_paragraph()
        linje.paragraph_format.left_indent = Pt(18)
        linje.paragraph_format.space_after = Pt(0)
        r = linje.add_run(f'{"*" if rigtigt else ""}{BOGSTAV[j]}. {svar}')
        if rigtigt:
            r.bold = True
            r.font.color.rgb = RGBColor(0x1A, 0x8F, 0x5E)
    nr = d.add_paragraph()
    nr.paragraph_format.left_indent = Pt(18)
    r = nr.add_run(f'Billede: {s["png"]} · {s["tid"]} sekunder')
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x8A, 0x93, 0xA6)

UD_DOCX = f'facit/forms-statistik-{DATO}.docx'
d.save(UD_DOCX)

# ------------------------------------------------------------ ren tekst
linjer = ['Statistik · 9. klasse · 30 spørgsmål',
          'Det rigtige svar er markeret med * foran.', '']
for s in SP:
    linjer.append(f'{s["nr"]}. {s["q"]}')
    for j, svar in enumerate(s['sv']):
        linjer.append(f'   {"*" if j == s["rigtig"] else " "}{BOGSTAV[j]}. {svar}')
    linjer.append('')
UD_TXT = f'facit/forms-statistik-{DATO}.txt'
open(UD_TXT, 'w').write('\n'.join(linjer))

# --------------------------------------------------------------- kontrol
kontrol = docx.Document(UD_DOCX)
stjerner = sum(1 for p in kontrol.paragraphs if p.text.lstrip().startswith('*'))
tekst = open(UD_TXT).read()
assert stjerner == len(SP), f'{stjerner} markerede facit, forventede {len(SP)}'
assert tekst.count('\n   *') == len(SP)
print(f'skrevet:  {UD_DOCX}   {len(SP)} spørgsmål · {stjerner} markerede facit')
print(f'skrevet:  {UD_TXT}')
